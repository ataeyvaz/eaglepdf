"""Yüksek sadakatli Word/Excel → PDF (Microsoft Office / LibreOffice): sahte motorla davranış testleri + makinede Office varsa gerçek deneme."""
import os
import subprocess
import sys

import pymupdf
import pytest

from eaglepdf.core import donustur, ofis
from eaglepdf.core.hatalar import BelgeHatasi

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Office/LibreOffice dönüştürme Windows'ta")


def ornek_pdf(yol):
    d = pymupdf.open()
    d.new_page().insert_text((72, 100), "Sahte motor çıktısı", fontname="helv")
    d.save(yol)
    d.close()
    return yol


def sahte_libreoffice(klasor, govde: str):
    """soffice.exe yerine geçen .cmd. Gövde, argümanlardan çözülen %OUT% (--outdir klasörü) ve %AD% (son argümanın adı, uzantısız)
    değişkenlerini kullanabilir (cmd `=` işaretinde argümanı böldüğü için konuma güvenilmez; bu yüzden döngüyle taranır)."""
    yol = klasor / "soffice.cmd"
    on = ('@echo off\r\n:dongu\r\nif "%~1"=="--outdir" set OUT=%~2\r\nset SON=%~1\r\nshift\r\n'
          'if not "%~1"=="" goto dongu\r\nfor %%F in ("%SON%") do set AD=%%~nF\r\n')
    yol.write_text(on + govde, encoding="utf-8")
    return str(yol)


def docx_yap(yol, metin="Merhaba dünya", ustbilgi="ÜSTBİLGİ ÖRNEĞİ"):
    import docx
    d = docx.Document()
    d.add_heading("Başlık ışık şğ", 1)
    d.add_paragraph(metin)
    d.sections[0].header.paragraphs[0].text = ustbilgi
    d.save(yol)
    return yol


def yalniz_libreoffice(monkeypatch, soffice: str):
    monkeypatch.setattr(ofis, "durum", lambda: ofis.OfisDurumu(False, False, soffice))


def test_motor_sirasi():
    d = ofis.OfisDurumu(True, True, r"C:\x\soffice.exe")
    assert d.motorlar("word") == ["Microsoft Word", "LibreOffice"]
    assert d.motorlar("excel") == ["Microsoft Excel", "LibreOffice"]
    assert ofis.OfisDurumu(False, False, None).motorlar("word") == []
    assert ofis.OfisDurumu(True, False, None).motorlar("excel") == []


def test_office_yoksa_yerlesik_donusturucuye_duser(tmp_path, monkeypatch):
    monkeypatch.setattr(ofis, "durum", lambda: ofis.OfisDurumu(False, False, None))
    docx_yap(tmp_path / "a.docx")
    pdf, ileti = donustur.ofisle_pdf("word", tmp_path / "a.docx", tmp_path / "o")
    assert "bulunamadı" in ileti and "Merhaba dünya" in pymupdf.open(pdf)[0].get_text()
    with pytest.raises(ofis.OfisYok):
        ofis.pdf_yap("word", tmp_path / "a.docx", tmp_path / "x.pdf")


def test_ofis_kapaliysa_hic_denenmez(tmp_path, monkeypatch):
    monkeypatch.setattr(ofis, "durum", lambda: pytest.fail("durum sorgulanmamalı"))
    docx_yap(tmp_path / "a.docx")
    pdf, ileti = donustur.ofisle_pdf("word", tmp_path / "a.docx", tmp_path / "o", ofis=False)
    assert ileti == "" and pdf.exists()


def test_libreoffice_ile_donusturur_turkce_ve_bosluklu_adla(tmp_path, monkeypatch):
    ornek = ornek_pdf(tmp_path / "hazir.pdf")
    monkeypatch.setenv("HAZIR_PDF", str(ornek))
    soffice = sahte_libreoffice(tmp_path, 'copy /Y "%HAZIR_PDF%" "%OUT%\\%AD%.pdf" >nul\r\n')
    yalniz_libreoffice(monkeypatch, soffice)
    kaynak = docx_yap(tmp_path / "Türkçe ad ışık.docx")
    pdf, ileti = donustur.ofisle_pdf("word", kaynak, tmp_path / "o")
    assert pdf.name == "Türkçe ad ışık.pdf" and "LibreOffice ile dönüştürüldü" in ileti
    assert "Sahte motor" in pymupdf.open(pdf)[0].get_text()
    assert kaynak.exists()                                          # kaynak dosya değişmedi


def test_motor_basarisizsa_yerlesige_duser_ve_nedenini_soyler(tmp_path, monkeypatch):
    yalniz_libreoffice(monkeypatch, sahte_libreoffice(tmp_path, "echo hata ayrintisi 1>&2\r\nexit /b 3\r\n"))
    kaynak = docx_yap(tmp_path / "a.docx")
    pdf, ileti = donustur.ofisle_pdf("word", kaynak, tmp_path / "o")
    assert "Office ile dönüştürülemedi" in ileti and "hata ayrintisi" in ileti and "Çıkış kodu 3" in ileti
    assert "Merhaba dünya" in pymupdf.open(pdf)[0].get_text()      # yerleşik çıktı üretildi
    assert [p.name for p in (tmp_path / "o").iterdir()] == ["a.pdf"]  # yarım Office çıktısı kalmadı


def test_zaman_asimi_sureci_oldurur(tmp_path, monkeypatch):
    yalniz_libreoffice(monkeypatch, sahte_libreoffice(tmp_path, "ping -n 60 127.0.0.1 >nul\r\n"))
    kaynak = docx_yap(tmp_path / "a.docx")
    with pytest.raises(BelgeHatasi) as h:
        ofis.pdf_yap("word", kaynak, tmp_path / "x.pdf", zaman_asimi=1.5)
    assert "LibreOffice" in str(h.value) and "bitmedi" in str(h.value)


def test_iptal_edilince_donusum_durur(tmp_path, monkeypatch):
    yalniz_libreoffice(monkeypatch, sahte_libreoffice(tmp_path, "ping -n 60 127.0.0.1 >nul\r\n"))
    kaynak = docx_yap(tmp_path / "a.docx")
    cagri = []

    def ilerleme(n, t):
        cagri.append(1)
        return len(cagri) > 3        # biraz bekledikten sonra iptal
    with pytest.raises(donustur.Iptal):
        donustur.ofisle_pdf("word", kaynak, tmp_path / "o", ilerleme)
    assert not list((tmp_path / "o").glob("*.pdf"))


def test_excel_yalniz_libreoffice_varsa_onu_kullanir(tmp_path, monkeypatch):
    import openpyxl
    ornek = ornek_pdf(tmp_path / "hazir.pdf")
    monkeypatch.setenv("HAZIR_PDF", str(ornek))
    yalniz_libreoffice(monkeypatch, sahte_libreoffice(tmp_path, 'copy /Y "%HAZIR_PDF%" "%OUT%\\%AD%.pdf" >nul\r\n'))
    wb = openpyxl.Workbook()
    wb.active.append([1, 2])
    wb.save(tmp_path / "t.xlsx")
    pdf, ileti = donustur.ofisle_pdf("excel", tmp_path / "t.xlsx", tmp_path / "o")
    assert "LibreOffice" in ileti and pdf.exists()


# --- gerçek Microsoft Office (makinede kuruluysa) ---
def _word_surecleri() -> set[str]:
    return ofis._surec_numaralari("WINWORD.EXE")


@pytest.mark.skipif(not ofis.durum().word, reason="Microsoft Word kurulu değil")
def test_gercek_word_ust_bilgiyi_korur_ve_sureci_kapatir(tmp_path):
    once = _word_surecleri()
    kaynak = docx_yap(tmp_path / "gerçek ad.docx")
    pdf, ileti = donustur.ofisle_pdf("word", kaynak, tmp_path / "o")
    assert "Microsoft Word ile dönüştürüldü" in ileti
    metin = pymupdf.open(pdf)[0].get_text()
    assert "ÜSTBİLGİ ÖRNEĞİ" in metin and "Merhaba dünya" in metin and "ışık" in metin    # yerleşik dönüştürücüde üstbilgi gelmez
    import time
    for _ in range(20):                      # Word'ün kapanması bir iki sn sürebilir
        if _word_surecleri() <= once:
            break
        time.sleep(0.5)
    assert _word_surecleri() <= once, "bizim başlattığımız Word süreci kapanmadı"


@pytest.mark.skipif(not ofis.durum().excel, reason="Microsoft Excel kurulu değil")
def test_gercek_excel_formulu_degeriyle_yazar(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapor"
    for i in range(1, 4):
        ws.append([i, i * 10, f"Kalem {i}"])
    ws["D1"] = "=SUM(B1:B3)"
    wb.save(tmp_path / "t.xlsx")
    pdf, ileti = donustur.ofisle_pdf("excel", tmp_path / "t.xlsx", tmp_path / "o")
    assert "Microsoft Excel ile dönüştürüldü" in ileti
    metin = pymupdf.open(pdf)[0].get_text()
    assert "Kalem 3" in metin and "60" in metin
