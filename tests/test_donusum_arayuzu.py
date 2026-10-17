import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import time
from pathlib import Path

import pymupdf
import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox

from eaglepdf import metin_tr
from eaglepdf.core import donustur, tarayici
from eaglepdf.core.belge import Belge, BelgeHatasi
from eaglepdf.ui import theme
from eaglepdf.ui.donusum_penceresi import TURLER, DonusumPenceresi, IsParcacigi
from eaglepdf.ui.karsilastirma_penceresi import KarsilastirmaPenceresi, fark_html, fark_ozeti
from eaglepdf.ui.pencere import AnaPencere
from test_belge import ARIAL, pdf_yap
from test_donustur import ornek_pdf, resim_olustur, taranmis_pdf


@pytest.fixture(scope="module")
def app():
    a = QApplication.instance() or QApplication([])
    theme.apply(a)
    return a


@pytest.fixture
def pencere(app, tmp_path):
    p = AnaPencere(str(tmp_path / "ayar.ini"))
    p.show()
    yield p
    if p.belge:
        p.belge.degisti = False
    p.close()


def bekle_kosul(fn, ms: int) -> bool:
    """`fn()` doğru olana ya da süre dolana dek olayları işleyerek bekler (PySide6'da QTest.qWaitFor yok)."""
    bitis = time.time() + ms / 1000
    while time.time() < bitis:
        QApplication.processEvents()
        if fn():
            return True
        time.sleep(0.01)
    return bool(fn())


def bitir(d: DonusumPenceresi, sure=60000):
    assert bekle_kosul(lambda: d.bitti, sure), "işlem zamanında bitmedi"


def sonuc_metinleri(d) -> list[str]:
    return [d.sonuclar.item(i).text() for i in range(d.sonuclar.count())]


def dialog(tur, dosyalar, cikti: Path):
    d = DonusumPenceresi(tur, [str(x) for x in dosyalar], str(cikti))
    return d


pytestmark = pytest.mark.skipif(not ARIAL.exists(), reason="Arial yazı tipi yok")


def test_turler_tam_ve_menu(pencere):
    assert set(TURLER) == {"pdf-word", "pdf-excel", "pdf-pptx", "pdf-resim", "pdf-metin", "pdf-html", "resim-pdf",
                           "word-pdf", "excel-pdf", "metin-pdf", "html-pdf", "sikistir"}
    m = pencere.m_donustur
    ust = [a.text() for a in m.actions() if a.text()]
    assert ust[:2] == ["PDF'den", "PDF'e"] and "PDF Sıkıştır…" in ust and "PDF Karşılaştır…" in ust and "Tarayıcıdan PDF Oluştur…" in ust
    assert [a.text() for a in m.actions()[0].menu().actions()] == [
        "Word (.docx)…", "Excel (.xlsx)…", "PowerPoint (.pptx)…", "Resim (PNG, JPG, BMP, TIFF)…", "Metin (.txt)…", "HTML (.html)…"]
    assert len(m.actions()[1].menu().actions()) == 5


def test_dosya_listesi(app, tmp_path):
    a = ornek_pdf(tmp_path / "a.pdf", 1)
    b = ornek_pdf(tmp_path / "b.pdf", 1)
    (tmp_path / "x.txt").write_text("x")
    d = dialog("pdf-metin", [a], tmp_path / "o")
    assert d.dosya_ekle(b) and not d.dosya_ekle(b) and not d.dosya_ekle(tmp_path / "x.txt") and not d.dosya_ekle(tmp_path / "yok.pdf")
    assert [x.name for x in d.dosyalar()] == ["a.pdf", "b.pdf"]
    d.liste.item(1).setSelected(True)
    d._tasi(-1)
    assert [x.name for x in d.dosyalar()] == ["b.pdf", "a.pdf"]
    d._kaldir()
    assert [x.name for x in d.dosyalar()] == ["a.pdf"] or [x.name for x in d.dosyalar()] == ["b.pdf"]
    bos = dialog("pdf-metin", [], tmp_path / "o")
    bos.basla()
    assert "en az bir dosya" in bos.durum.text() and bos.is_parcacigi is None


def test_pdf_word_toplu_ve_hata(app, tmp_path):
    a = ornek_pdf(tmp_path / "a.pdf", 1)
    t = taranmis_pdf(tmp_path / "tara.pdf")
    cikti = tmp_path / "o"
    d = dialog("pdf-word", [a, t], cikti)
    d.basla()
    bitir(d)
    s = sonuc_metinleri(d)
    assert s[0].startswith("✓ a.pdf") and s[1].startswith("✗ tara.pdf") and "OCR" in s[1]
    assert "1/2 başarılı" in d.durum.text()
    assert (cikti / "a.docx").is_file() and d.ciktilar == [str(cikti / "a.docx")]
    assert d.acilacak == d.ciktilar and not d.tur.cikti_pdf   # PDF olmayan çıktıda pencere klasörü gösterir
    assert d.baslat_dugme.text() == "Yeniden Dönüştür"


def test_secenekli_donusumler(app, tmp_path):
    a = ornek_pdf(tmp_path / "a.pdf", 3)
    o = tmp_path / "o"
    d = dialog("pdf-resim", [a], o)
    d.sayfalar_kutu.setText("2-3")
    d.bicim_kutu.setCurrentText("jpg")
    d.dpi_kutu.setValue(72)
    d.renk_kutu.setCurrentIndex(1)
    d.basla()
    bitir(d)
    assert sorted(p.name for p in o.glob("*")) == ["a_sayfa_002.jpg", "a_sayfa_003.jpg"] and "2 resim" in sonuc_metinleri(d)[0]

    d = dialog("pdf-excel", [a], o)
    d.basla()
    bitir(d)
    assert "3 tablo" in sonuc_metinleri(d)[0]

    d = dialog("pdf-pptx", [a], o)
    d.pptx_mod.setCurrentIndex(1)
    d.sayfalar_kutu.setText("1")
    d.basla()
    bitir(d)
    from pptx import Presentation
    assert len(Presentation(str(o / "a.pptx")).slides) == 1

    d = dialog("pdf-metin", [a], o)
    d.ayrac_kutu.setChecked(False)
    d.basla()
    bitir(d)
    assert "Sayfa 1 —" not in (o / "a.txt").read_text(encoding="utf-8-sig")

    d = dialog("pdf-html", [a], o)
    d.basla()
    bitir(d)
    assert (o / "a.html").is_file()


def test_gecersiz_sayfa_araligi_ve_parola(app, tmp_path):
    a = ornek_pdf(tmp_path / "a.pdf", 2)
    d = dialog("pdf-metin", [a], tmp_path / "o")
    d.sayfalar_kutu.setText("9")
    d.basla()
    assert d.is_parcacigi is None and "1–2" in d.durum.text()
    p = pdf_yap(tmp_path / "p.pdf", 1, "Gizli metin", parola="abc")
    d = dialog("pdf-metin", [p], tmp_path / "o")
    d.basla()
    bitir(d)
    assert sonuc_metinleri(d)[0].startswith("✗") and "parola" in sonuc_metinleri(d)[0].lower()
    d = dialog("pdf-metin", [p], tmp_path / "o")
    d.parola_kutu.setText("abc")
    d.basla()
    bitir(d)
    assert sonuc_metinleri(d)[0].startswith("✓")


def test_resimlerden_pdf_penceresi(app, tmp_path):
    r1 = resim_olustur(tmp_path / "1.png", (300, 200))
    r2 = resim_olustur(tmp_path / "2.png", (200, 300))
    d = dialog("resim-pdf", [r1, r2], tmp_path / "o")
    d.ad_kutu.setText("Albüm")
    d.kenar_kutu.setValue(10)
    d.basla()
    bitir(d)
    pdf = tmp_path / "o" / "Albüm.pdf"
    assert pymupdf.open(pdf).page_count == 2 and d.acilacak == [str(pdf)]
    assert d.sonuclar.count() == 1 and "2 resim" in sonuc_metinleri(d)[0]
    d2 = dialog("resim-pdf", [r1, r2], tmp_path / "o2")
    d2.ayri_kutu.setChecked(True)
    d2.basla()
    bitir(d2)
    assert sorted(p.name for p in (tmp_path / "o2").glob("*.pdf")) == ["1.pdf", "2.pdf"]


def test_diger_pdfe_donusumler(app, tmp_path):
    import docx
    import openpyxl
    dd = docx.Document()
    dd.add_paragraph("Word içeriği Işık")
    dd.save(tmp_path / "a.docx")
    wb = openpyxl.Workbook()
    wb.active.append(["Çay", 12])
    wb.save(tmp_path / "a.xlsx")
    (tmp_path / "a.txt").write_text("Düz metin ğüş", encoding="utf-8")
    (tmp_path / "a.html").write_text("<html><body><h1>Başlık</h1></body></html>", encoding="utf-8")
    for tur, ad, aranan in (("word-pdf", "a.docx", "Word içeriği Işık"), ("excel-pdf", "a.xlsx", "Çay"),
                            ("metin-pdf", "a.txt", "Düz metin ğüş"), ("html-pdf", "a.html", "Başlık")):
        d = dialog(tur, [tmp_path / ad], tmp_path / f"o-{tur}")
        if hasattr(d, "ofis_kutu"):
            d.ofis_kutu.setChecked(False)   # bu test yerleşik dönüştürücüyü sınar (makinede Office olsa da değişmesin)
        d.basla()
        bitir(d)
        assert sonuc_metinleri(d)[0].startswith("✓"), (tur, sonuc_metinleri(d))
        pdf = tmp_path / f"o-{tur}" / "a.pdf"
        assert aranan in pymupdf.open(pdf)[0].get_text() and d.acilacak == [str(pdf)]


def test_sikistir_penceresi(app, tmp_path):
    from test_donustur import buyuk_resimli_pdf
    k = buyuk_resimli_pdf(tmp_path / "foto.pdf")
    d = dialog("sikistir", [k], tmp_path / "o")
    assert not d.ac_kutu.isChecked()  # sıkıştırmada sonuç varsayılan olarak açılmaz
    d.seviye_kutu.setCurrentIndex(2)
    d.basla()
    bitir(d)
    m = sonuc_metinleri(d)[0]
    assert m.startswith("✓ foto.pdf") and "küçüldü" in m and "MB" in m
    assert (tmp_path / "o" / "foto_sikistirilmis.pdf").stat().st_size < k.stat().st_size * 0.3 and d.acilacak == []


def test_iptal(app, tmp_path):
    adimlar = []

    def uzun(cb):
        for n in range(200):
            adimlar.append(n)
            if cb(n, 200):
                raise donustur.Iptal("iptal")
            time.sleep(0.01)
        return [], "bitti"
    ikinci = []
    ip = IsParcacigi([("uzun", uzun), ("ikinci", lambda cb: (ikinci.append(1), ([], ""))[1])])
    sonuclar = []
    ip.dosyaBitti.connect(lambda i, ok, m, c: sonuclar.append((i, ok, m)))
    ip.start()
    QTest.qWait(80)
    ip.iptal()
    assert ip.wait(5000)
    QApplication.processEvents()
    assert sonuclar == [(0, False, "İptal edildi.")] and ikinci == [] and len(adimlar) < 200


def test_arayuzden_donustur_sonrasi_pdf_acilir(pencere, tmp_path, app):
    r = resim_olustur(tmp_path / "r.png")
    d = dialog("resim-pdf", [r], tmp_path / "o")
    d.basla()
    bitir(d)
    pencere._donusum_sonrasi(d)
    bekle = bekle_kosul(lambda: pencere.belge is not None, 5000)
    assert bekle and pencere.belge.yol == tmp_path / "o" / "r.pdf"


def test_kayitli_pdf_yolu_degismis_belge(pencere, tmp_path, app, monkeypatch):
    pencere.ac(ornek_pdf(tmp_path / "a.pdf", 1))
    assert pencere._kayitli_pdf_yolu() == str(tmp_path / "a.pdf")
    pencere.belge.ek_not(0, (100, 100), "x")
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Cancel)
    assert pencere._kayitli_pdf_yolu() == ""
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.No)
    assert pencere._kayitli_pdf_yolu() == str(tmp_path / "a.pdf")
    kaydedilen = []
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    monkeypatch.setattr(pencere, "farkli_kaydet", lambda: kaydedilen.append(1) or True)
    assert pencere._kayitli_pdf_yolu() == str(tmp_path / "a.pdf") and kaydedilen == [1]
    monkeypatch.setattr(pencere, "farkli_kaydet", lambda: False)
    assert pencere._kayitli_pdf_yolu() == ""


# --- karşılaştırma penceresi ---
def test_karsilastirma_penceresi(app, tmp_path):
    a = pdf_yap(tmp_path / "a.pdf", 3, "Merhaba dunya")
    c = pdf_yap(tmp_path / "c.pdf", 3, "Merhaba dunya")
    d = pymupdf.open(c)
    d[1].insert_text((72, 200), "Eklenen cümle burada", fontname="helv", fontsize=14)
    d.saveIncr()
    d.close()
    w = KarsilastirmaPenceresi(str(a), str(c))
    w.basla()
    assert bekle_kosul(lambda: w.liste.count() == 3, 30000)
    assert "1 sayfada fark var" in w.ozet.text()
    assert [w.liste.item(i).text().split(" — ")[1] for i in range(3)][0] == "fark yok"
    assert w.liste.currentRow() == 1                      # ilk farklı sayfa seçilir
    assert not w.gorsel.pixmap().isNull() and "Eklenen cümle burada" in w.metin.toPlainText()
    w.liste.setCurrentRow(0)
    assert w.gorsel.pixmap().isNull() and "yok" in w.gorsel.text()
    # kontroller
    w2 = KarsilastirmaPenceresi(str(a), str(a))
    w2.basla()
    assert "aynı dosyayı" in metin_tr.kucuk(w2.ozet.text()) and w2.is_parcacigi is None
    w3 = KarsilastirmaPenceresi(str(a), "")
    w3.basla()
    assert "iki belgeyi" in metin_tr.kucuk(w3.ozet.text())
    w4 = KarsilastirmaPenceresi(str(a), str(a.with_name("b.pdf")))
    (tmp_path / "b.pdf").write_bytes(b"degil")
    w4.basla()
    assert bekle_kosul(lambda: w4.ozet.text().startswith("⚠"), 15000)


def test_fark_metinleri():
    f = donustur.SayfaFarki(0, None, ["yeni <b>"], ["eski"], 0.02, None)
    assert "metin değişti" in fark_ozeti(f) and "%2,0" in fark_ozeti(f)
    h = fark_html(f)
    assert "yeni &lt;b&gt;" in h and "eski" in h and "#2e7d32" in h
    assert fark_ozeti(donustur.SayfaFarki(1, "b", [], [], 1.0, None)) == "yalnızca ikinci belgede var"
    assert fark_ozeti(donustur.SayfaFarki(1, None, [], [], 0.0, None)) == "fark yok"
    assert "kelime farkı yok" in fark_html(donustur.SayfaFarki(1, None, [], [], 0.1, None))


# --- tarayıcı ---
class SahteResim:
    def __init__(self, kaynak):
        self.kaynak = kaynak

    def SaveFile(self, yol):
        Path(yol).write_bytes(Path(self.kaynak).read_bytes())


class SahteDiyalog:
    def __init__(self, sonuc):
        self.sonuc = sonuc

    def ShowAcquireImage(self, *a):
        if isinstance(self.sonuc, Exception):
            raise self.sonuc
        return self.sonuc


class SahteWiaHatasi(Exception):
    def __init__(self, hresult):
        super().__init__(f"hata {hresult:#x}")
        self.hresult = hresult


def test_tarayici_cekirdek(tmp_path):
    png = resim_olustur(tmp_path / "kaynak.png")
    yol = tarayici.tara(tmp_path / "t", "gri", lambda ad: SahteDiyalog(SahteResim(png)))
    assert yol.suffix == ".png" and yol.read_bytes() == png.read_bytes()
    with pytest.raises(donustur.Iptal):
        tarayici.tara(tmp_path / "t", "renkli", lambda ad: SahteDiyalog(None))
    with pytest.raises(donustur.Iptal):
        tarayici.tara(tmp_path / "t", "renkli", lambda ad: SahteDiyalog(SahteWiaHatasi(tarayici.HATA_KULLANICI_IPTAL)))
    with pytest.raises(BelgeHatasi, match="Tarayıcı bulunamadı"):
        tarayici.tara(tmp_path / "t", "renkli", lambda ad: SahteDiyalog(SahteWiaHatasi(tarayici.HATA_CIHAZ_YOK)))
    with pytest.raises(BelgeHatasi, match="kâğıt"):
        tarayici.tara(tmp_path / "t", "renkli", lambda ad: SahteDiyalog(SahteWiaHatasi(tarayici.HATA_KAGIT_YOK)))
    with pytest.raises(BelgeHatasi, match="Tarama başarısız"):
        tarayici.tara(tmp_path / "t", "renkli", lambda ad: SahteDiyalog(SahteWiaHatasi(0x80004005)))
    with pytest.raises(BelgeHatasi):
        tarayici.tara(tmp_path / "t", "mor", lambda ad: SahteDiyalog(None))
    assert isinstance(tarayici.tarayicilar(), list)  # cihaz yoksa boş liste, hata yok


def test_tarayicidan_pdf_arayuzu(pencere, tmp_path, app, monkeypatch):
    png = resim_olustur(tmp_path / "s.png", (200, 300))
    sayac = []

    def sahte_tara(klasor, renk):
        sayac.append(renk)
        kopya = Path(klasor) / f"s{len(sayac)}.png"
        kopya.write_bytes(png.read_bytes())
        return kopya
    monkeypatch.setattr(tarayici, "kullanilabilir", lambda: True)
    monkeypatch.setattr(tarayici, "tara", sahte_tara)
    monkeypatch.setattr(QInputDialog, "getItem", lambda *a, **k: ("Gri tonlama", True))
    cevaplar = iter([QMessageBox.Yes, QMessageBox.No])
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: next(cevaplar))
    hedef = tmp_path / "Tarama.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(hedef), ""))
    pencere.tarayicidan_pdf()
    assert sayac == ["gri", "gri"] and pymupdf.open(hedef).page_count == 2
    assert pencere.belge.yol == hedef
    # iptal: dosya oluşmaz
    pencere.belge.degisti = False

    def iptal_tara(klasor, renk):
        raise donustur.Iptal("x")
    monkeypatch.setattr(tarayici, "tara", iptal_tara)
    hedef2 = tmp_path / "Iptal.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(hedef2), ""))
    pencere.tarayicidan_pdf()
    assert not hedef2.exists()
    # hata: mesaj gösterilir
    gosterilen = []
    monkeypatch.setattr(tarayici, "tara", lambda k, r: (_ for _ in ()).throw(BelgeHatasi("Tarayıcı bulunamadı.")))
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: gosterilen.append(a[2]))
    pencere.tarayicidan_pdf()
    assert gosterilen == ["Tarayıcı bulunamadı."]
    # WIA yoksa bilgi verilir
    monkeypatch.setattr(tarayici, "kullanilabilir", lambda: False)
    bilgi = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: bilgi.append(a[2]))
    pencere.tarayicidan_pdf()
    assert bilgi and "Windows" in bilgi[0]


def test_word_excel_pdf_ofis_secenegi(app, tmp_path, monkeypatch):
    from eaglepdf.core import ofis
    monkeypatch.setattr(ofis, "durum", lambda: ofis.OfisDurumu(False, False, None))
    d = dialog("word-pdf", [], tmp_path / "o")
    assert not d.ofis_kutu.isEnabled() and not d.ofis_kutu.isChecked()          # Office yok: seçenek kapalı
    assert d._secenek_degerleri() == {"ofis": False}
    monkeypatch.setattr(ofis, "durum", lambda: ofis.OfisDurumu(True, True, None))
    w = dialog("word-pdf", [], tmp_path / "o")
    assert w.ofis_kutu.isEnabled() and w.ofis_kutu.isChecked() and w._secenek_degerleri() == {"ofis": True}
    e = dialog("excel-pdf", [], tmp_path / "o")
    from PySide6.QtWidgets import QLabel
    assert e.ofis_kutu.isEnabled() and any("Bulundu: Microsoft Excel" in e_.text() for e_ in e.secenekler_kutusu.findChildren(QLabel))
