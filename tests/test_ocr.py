import pymupdf
import pytest

from eaglepdf import metin_tr
from eaglepdf.core import ocr, yazitipi
from eaglepdf.core.belge import Belge
from eaglepdf.core.hatalar import BelgeHatasi

pytestmark = pytest.mark.skipif(not ocr.hazir_mi(), reason="Türkçe+İngilizce OCR dil verisi kurulu değil")

CUMLE = "Işık İstanbul şğüöç"


def taranmis(yol, metinler=(CUMLE,), ekstra_metinli=False):
    """Metni resme çevrilmiş (metin katmanı olmayan) PDF üretir."""
    kaynak = pymupdf.open()
    for m in metinler:
        s = kaynak.new_page()
        s.insert_font(fontname="noto", fontbuffer=yazitipi.tampon())
        s.insert_text((72, 120), m, fontsize=26, fontname="noto")
    hedef = pymupdf.open()
    for s in kaynak:
        yeni = hedef.new_page(width=s.rect.width, height=s.rect.height)
        yeni.insert_image(yeni.rect, pixmap=s.get_pixmap(dpi=200))
    if ekstra_metinli:
        s = hedef.new_page()
        s.insert_text((72, 100), "Zaten metinli", fontname="helv")
    hedef.save(yol)
    return Belge.ac(yol)


def test_dil_verisi_bulunur():
    assert "tur" in ocr.kullanilabilir_diller()
    assert ocr.hazir_mi("tur")
    assert not ocr.hazir_mi("xyz")


def test_aranabilir_yap_turkce(tmp_path):
    b = taranmis(tmp_path / "t.pdf")
    assert b.metinsiz_sayfalar() == [0]
    sonuc = b.aranabilir_yap()
    assert sonuc.islenen == 1 and sonuc.kelime >= 3
    metin = b._doc[0].get_text()
    for kelime in ("Işık", "İstanbul", "şğüöç"):
        assert kelime in metin, metin
    # arama (Türkçe büyük/küçük harf duyarlı yardımcı ile)
    assert b.ara("istanbul")
    assert metin_tr.kucuk("İSTANBUL") == "istanbul"


def test_gorunum_degismez_ve_kaydedilir(tmp_path):
    b = taranmis(tmp_path / "t.pdf")
    once = b._doc[0].get_pixmap(dpi=50).samples
    b.aranabilir_yap()
    assert b._doc[0].get_pixmap(dpi=50).samples == once  # metin görünmez
    b.kaydet(tmp_path / "ocr.pdf")
    doc = pymupdf.open(tmp_path / "ocr.pdf")
    assert "stanbul" in doc[0].get_text()
    doc.close()


def test_metinli_sayfa_atlanir(tmp_path):
    b = taranmis(tmp_path / "t.pdf", ekstra_metinli=True)
    sonuc = b.aranabilir_yap()
    assert (sonuc.islenen, sonuc.atlanan_metinli) == (1, 1)


def test_geri_al(tmp_path):
    b = taranmis(tmp_path / "t.pdf")
    b.aranabilir_yap()
    assert b._doc[0].get_text().strip()
    assert b.geri_al()
    assert not b._doc[0].get_text().strip()


def test_iptal_belgeyi_degistirmez(tmp_path):
    b = taranmis(tmp_path / "t.pdf", metinler=(CUMLE, CUMLE))
    with pytest.raises(ocr.OcrIptal):
        b.aranabilir_yap(ilerleme=lambda n, t: n >= 1)
    assert b.metinsiz_sayfalar() == [0, 1]
    assert not b.degisti


def test_bolgeden_metin(tmp_path):
    b = taranmis(tmp_path / "t.pdf")
    metin = b.bolgeden_metin(0, (50, 90, 545, 140))
    assert "stanbul" in metin
    assert not b.degisti


def test_dil_verisi_yoksa_anlasilir_hata(tmp_path):
    b = taranmis(tmp_path / "t.pdf")
    with pytest.raises(BelgeHatasi, match="dil verisi"):
        b.aranabilir_yap(dil="xyz")


# --- Faz 4 dönüşümlerinin OCR'a bağlanması ---
def test_donusumler_taranmis_belgede_ocr_kullanir(tmp_path):
    import docx
    from eaglepdf.core import donustur
    kaynak = tmp_path / "tara.pdf"
    taranmis(kaynak).kapat()
    once = kaynak.read_bytes()
    txt = donustur.pdf_metin(kaynak, tmp_path / "m").read_text(encoding="utf-8-sig")
    assert "stanbul" in txt
    assert "stanbul" in donustur.pdf_html(kaynak, tmp_path / "h").read_text(encoding="utf-8")
    d = docx.Document(str(donustur.pdf_word(kaynak, tmp_path / "w")))
    assert "stanbul" in "\n".join(p.text for p in d.paragraphs)
    assert kaynak.read_bytes() == once  # kaynak dosya değişmez


def test_pptx_taranmis_sayfa_notlara_ocr_metni(tmp_path):
    from pptx import Presentation
    from eaglepdf.core import donustur
    kaynak = tmp_path / "tara.pdf"
    taranmis(kaynak).kapat()
    for mod in ("duzenlenebilir", "resim"):
        dosya = donustur.pdf_pptx(kaynak, tmp_path / mod, mod=mod)
        slayt = Presentation(str(dosya)).slides[0]
        assert "stanbul" in slayt.notes_slide.notes_text_frame.text
        kutular = [sh for sh in slayt.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        assert kutular == []  # resimle çakışan ikinci metin katmanı yok
